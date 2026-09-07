#!/usr/bin/env python3
import argparse
import re
from typing import Optional

from openpyxl import load_workbook


def normalize_col_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def find_header(ws, preferred_names):
    headers = [ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)]
    for idx, header in enumerate(headers, start=1):
        if header is None:
            continue
        for name in preferred_names:
            if normalize_col_name(str(header)) == normalize_col_name(name):
                return idx
    return None


def parse_number(value):
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace(".", "").replace(",", ".")
        try:
            return float(text)
        except ValueError:
            m = re.search(r"[-+]?\d+(?:[\.,]\d+)?", text)
            if m:
                return float(m.group(0).replace(".", "").replace(",", "."))
    return 0.0


def format_tl(value: float) -> str:
    return f"{value:.1f}".replace(".", ",")


def parse_cost_tl(value) -> tuple[Optional[float], Optional[float]]:
    if not isinstance(value, str):
        return None, None

    text = value.strip()
    m = re.search(r"([0-9,\.]+)\s*USD", text, re.IGNORECASE)
    if not m:
        return None, None

    usd = float(m.group(1).replace(".", "").replace(",", "."))

    m_tl = re.search(r"=\s*([0-9,\.]+)\s*(?:TL)?", text, re.IGNORECASE)
    tl = None
    if m_tl:
        tl = float(m_tl.group(1).replace(".", "").replace(",", "."))

    return usd, tl


def format_cost_text(usd_value: float, tl_value: float) -> str:
    usd_text = f"{usd_value:.2f}".replace(".", ",")
    tl_text = format_tl(tl_value)
    return f"{usd_text} USD = {tl_text}"


def update_workbook(file_path: str, rate: float, sheet_name: Optional[str] = None) -> int:
    wb = load_workbook(file_path)
    ws = wb[sheet_name] if sheet_name else wb.active

    cost_col = find_header(ws, ["MALİYET", "Maliyet", "COST"])
    commission_col = find_header(ws, ["KOMİSYON %", "Komisyon %", "Commission %"])
    kargo_col = find_header(ws, ["KARGO", "Kargo"])
    platform_col = find_header(ws, ["PLATFORM", "Platform"])
    paketleme_col = find_header(ws, ["PAKETLEME", "Paketleme"])
    patpat_col = find_header(ws, ["PATPAT", "Patpat"])
    bant_col = find_header(ws, ["BANT", "Bant"])
    sale_col = find_header(ws, ["GÜNCEL SATIŞI", "Guncel Satisi", "Current Sale"])
    rent_col = find_header(ws, ["KİRA- ELEKTRİK", "Kira-Elektrik", "Kira Elektrik"])
    profit_col = find_header(ws, ["KAR MARJI", "Kar Marjı", "Profit"])

    if None in [cost_col, commission_col, sale_col, profit_col]:
        raise ValueError("Gerekli sütunlar bulunamadı. Dosya yapısını kontrol edin.")

    updated = 0
    for row in range(2, ws.max_row + 1):
        cost_cell = ws.cell(row=row, column=cost_col)
        commission_cell = ws.cell(row=row, column=commission_col)
        sale_cell = ws.cell(row=row, column=sale_col)
        profit_cell = ws.cell(row=row, column=profit_col)

        usd, _ = parse_cost_tl(cost_cell.value)
        if usd is None:
            continue

        new_tl = round(usd * rate, 1)
        cost_cell.value = format_cost_text(usd, new_tl)

        commission_pct = parse_number(commission_cell.value)
        if commission_pct > 1:
            commission_pct = commission_pct / 100.0

        fixed_costs = 0.0
        for col in [kargo_col, platform_col, paketleme_col, patpat_col, bant_col, rent_col]:
            if col is not None:
                fixed_costs += parse_number(ws.cell(row=row, column=col).value)

        current_profit = parse_number(profit_cell.value)
        if current_profit == 0.0:
            current_profit = parse_number(sale_cell.value) - (new_tl + fixed_costs)

        sale_value = (new_tl + fixed_costs + current_profit) / (1 - commission_pct) if commission_pct < 1 else 0.0
        sale_cell.value = round(sale_value, 1)
        profit_cell.value = round(sale_value - (new_tl + fixed_costs + sale_value * commission_pct), 1)
        updated += 1

    wb.save(file_path)
    return updated


def main():
    parser = argparse.ArgumentParser(description="MASIS workbook’ını USD/TRY kuruna göre günceller")
    parser.add_argument("--file", required=True, help="Excel dosyasının yolu")
    parser.add_argument("--rate", type=float, required=True, help="USD/TRY kuru")
    parser.add_argument("--sheet", default=None, help="İşlenecek sayfa adı")
    args = parser.parse_args()

    updated_count = update_workbook(args.file, args.rate, args.sheet)
    print(f"Güncellendi: {updated_count} satır")


if __name__ == "__main__":
    main()
