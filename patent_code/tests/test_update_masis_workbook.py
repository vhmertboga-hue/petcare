import os
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from update_masis_workbook import update_workbook


class UpdateWorkbookTests(unittest.TestCase):
    def test_updates_cost_and_profit_without_tl_suffix(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Sayfa1"
        ws.append(["ÜRÜN", "MALİYET", "KOMİSYON %", "KARGO", "PLATFORM", "PAKETLEME", "PATPAT ", "BANT", "GÜNCEL SATIŞI", "KİRA- ELEKTRİK", "KAR MARJI"])
        ws.append(["MS-01", "13,33 USD", 20, 93, 12.5, 10, 7.5, 2.5, 1000, 25, 200])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = os.path.join(tmpdir, "test.xlsx")
            wb.save(path)

            updated = update_workbook(path, rate=47.5)

            self.assertEqual(updated, 1)
            ws2 = load_workbook(path).active
            self.assertEqual(ws2["B2"].value, "13,33 USD = 633,2")
            self.assertIsNotNone(ws2["I2"].value)
            self.assertIsNotNone(ws2["K2"].value)


if __name__ == "__main__":
    unittest.main()
